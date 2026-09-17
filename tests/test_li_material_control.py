from pathlib import Path

from openpyxl import Workbook, load_workbook

from pipeline import li_material_control


def _make_baseline(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "1 pav"
    ws.column_dimensions["G"].width = 12
    ws.column_dimensions["H"].width = 12
    ws["G5"] = "REV. 0"
    ws["G6"] = 10
    ws["G7"] = "A LEVANTAR"
    ws.print_area = "A1:J20"
    wb.save(path)


def _make_current(baseline: Path, current: Path, change_old: bool = False) -> None:
    wb = load_workbook(baseline)
    ws = wb["1 pav"]
    ws["H5"] = "REV. A"
    ws["H6"] = 15
    ws["H7"] = 8
    if change_old:
        ws["G6"] = 99
    wb.save(current)


def test_new_revision_preserves_rev0(tmp_path: Path) -> None:
    baseline = tmp_path / "baseline.xlsx"
    current = tmp_path / "current.xlsx"
    _make_baseline(baseline)
    _make_current(baseline, current)
    result = li_material_control.validate(current, baseline, "REV. 0", "REV. A")
    assert result["status"] == "VALIDADO"
    assert result["findings"] == []


def test_historical_revision_change_is_reproved(tmp_path: Path) -> None:
    baseline = tmp_path / "baseline.xlsx"
    current = tmp_path / "current.xlsx"
    _make_baseline(baseline)
    _make_current(baseline, current, change_old=True)
    result = li_material_control.validate(current, baseline, "REV. 0", "REV. A")
    assert result["status"] == "REPROVADO"
    assert any(x.startswith("historical_revision_changed:") for x in result["findings"])


def test_missing_baseline_is_hold(tmp_path: Path) -> None:
    baseline = tmp_path / "baseline.xlsx"
    current = tmp_path / "current.xlsx"
    _make_baseline(baseline)
    _make_current(baseline, current)
    result = li_material_control.validate(current, None, "REV. 0", "REV. A")
    assert result["status"] == "HOLD"
    assert "baseline_missing:historical_preservation_not_proven" in result["findings"]
