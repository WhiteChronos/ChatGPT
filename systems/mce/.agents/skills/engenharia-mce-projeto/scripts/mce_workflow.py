#!/usr/bin/env python3
"""MCE: workflow verificavel, local e com aprovacao humana obrigatoria.

Fase 1: init -> understand -> STOP. Fase 2 somente: authorize -> apply -> audit.
Nunca modifica o XLSX mestre; edicao cirurgica de XML dentro de copia ZIP.
"""
from __future__ import annotations

import argparse
import copy
import datetime as dt
import hashlib
import json
import posixpath
import re
import shutil
import sys
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.cell import coordinate_from_string

MAIN = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
RNS = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
PKG = 'http://schemas.openxmlformats.org/package/2006/relationships'
NS = {'m': MAIN, 'r': RNS}
M = lambda name: '{'+MAIN+'}'+name
GRAY = 'FFD9D9D9'
PHASES = ('INVENTARIADO', 'ANALISADO_AGUARDANDO_DECISAO', 'AUTORIZADO', 'MCE_PRELIMINAR_PARA_REVISAO')


def utc():
    return dt.datetime.now(dt.timezone.utc).isoformat(timespec='seconds')


def digest(p):
    h = hashlib.sha256()
    with Path(p).open('rb') as f:
        for block in iter(lambda:f.read(1048576), b''):
            h.update(block)
    return h.hexdigest()


def readjson(path):
    return json.loads(Path(path).read_text(encoding='utf-8'))


def writejson(path, data):
    p=Path(path); p.parent.mkdir(parents=True,exist_ok=True)
    p.write_text(json.dumps(data,indent=2,ensure_ascii=False,sort_keys=True)+'\n',encoding='utf-8')


def fail(msg):
    raise ValueError(msg)


def master_path(ws):
    return Path(ws)/'base'/'modelo.xlsx'


def load_state(ws):
    p=Path(ws)/'state.json'
    if not p.exists(): fail('Workspace nao inicializado; execute init.')
    s=readjson(p)
    if digest(master_path(ws)) != s['template_sha256']:
        fail('STOP: hash da planilha mestre divergente do inventario.')
    for doc in s['sources']:
        p=Path(ws)/'sources'/doc['name']
        if not p.is_file() or digest(p)!=doc['sha256']:
            fail('STOP: fonte alterada ou ausente: '+doc['name'])
    return s


def count_formula_cells(sheet):
    return sum(isinstance(c.value,str) and c.value.startswith('=') for row in sheet for c in row)


def inventory(p):
    w=load_workbook(p,read_only=False,keep_links=True)
    return {s.title:{'rows':s.max_row,'columns':s.max_column,
                     'formulas':count_formula_cells(s),
                     'merged':sorted(str(x) for x in s.merged_cells.ranges),
                     'print_area':str(s.print_area or ''),
                     'protected':bool(s.protection.sheet)} for s in w}


def run_init(args):
    w=Path(args.workspace).resolve()
    if (w/'state.json').exists():fail('Workspace existente; proibido sobrescrever baseline.')
    src=Path(args.master).resolve()
    if src.suffix.lower()!='.xlsx':fail('Apenas .xlsx, sem macros, e suportado pelo editor atual.')
    if not src.is_file():fail('Modelo Excel nao encontrado.')
    w.mkdir(parents=True,exist_ok=True); (w/'base').mkdir(exist_ok=True)
    (w/'sources').mkdir(exist_ok=True); (w/'output').mkdir(exist_ok=True)
    shutil.copy2(src, master_path(w))
    docs=[]; seen=set()
    for p0 in args.sources:
        p=Path(p0).resolve()
        if not p.is_file():fail('Documento ausente: '+str(p))
        if p.name in seen:fail('Nomes de documentos duplicados; diferencie as revisoes: '+p.name)
        seen.add(p.name)
        shutil.copy2(p,w/'sources'/p.name)
        docs.append({'name':p.name,'sha256':digest(p),'size_bytes':p.stat().st_size,
                     'extension':p.suffix.lower(),'revision':'NAO INFORMADA'})
    s={'version':1,'status':'INVENTARIADO','created_at':utc(),
       'template_name':src.name,'template_sha256':digest(master_path(w)),
       'template_classification':args.classification,'sources':docs,
       'source_manifest_sha256':hashlib.sha256(json.dumps(docs,sort_keys=True).encode()).hexdigest(),
       'workbook_inventory':inventory(master_path(w)), 'approval':None,
       'artifacts':[]}
    writejson(w/'state.json',s)
    writejson(w/'contract.json',{'template_sha256':s['template_sha256'],
        'model_approved':False,'max_causes_per_sheet':50,
        'matrix_ranges':[],'editable_cells':{},
        'protected_rows_by_sheet':{},'frozen_sheets':[]})
    print(json.dumps({'status':s['status'],'workspace':str(w),'master_sha256':s['template_sha256'],
          'documents':len(docs)},ensure_ascii=False))


def validate_finding(ws,state,item):
    if item['classification'] not in ('DOCUMENTADO','INFERIDO_PARA_REVISAO','NAO_DOCUMENTADO'):
        fail('Classificacao invalida em analise')
    if item['classification']=='DOCUMENTADO':
        doc=item.get('source_name'); excerpt=item.get('verbatim_excerpt')
        if not doc or not excerpt:fail('Fato documentado deve conter fonte e trecho literal.')
        f=next((x for x in state['sources'] if x['name']==doc),None)
        if not f:fail('Documento nao inventariado: '+str(doc))
        if f['extension'] in ('.txt','.md'):
            lines=(Path(ws)/'sources'/doc).read_text(encoding='utf-8-sig').splitlines()
            loc=item.get('line')
            if not isinstance(loc,int) or loc<1 or loc>len(lines) or excerpt not in lines[loc-1]:
                fail('Trecho nao encontrado na linha declarada em '+doc)
        elif item.get('manual_evidence_verified') is not True:
            fail('Evidencia de arquivo binario requer revisao manual independente.')


def run_understand(args):
    w=Path(args.workspace);s=load_state(w)
    if s['status']!='INVENTARIADO':fail('Analise ja registrada; criar nova revisao em outro workspace.')
    a=readjson(args.analysis)
    if not a.get('summary') or not isinstance(a.get('findings'),list) or not isinstance(a.get('gaps'),list):
        fail('Analise exige resumo, fatos e pendencias.')
    # Evidencias de P&ID/OCR/benchmark nao se tornam conclusoes documentadas
    # apenas porque um adaptador produziu um JSON. Exigir revisao humana por SHA.
    derived=s.get('pid_integrations',[])
    if derived:
        reviews=a.get('pid_integration_review',[])
        if not isinstance(reviews,list) or len(reviews)!=len(derived):
            fail('STOP: toda integracao P&ID precisa de revisao humana no analise.json.')
        refs={x['output_file']:x for x in derived}
        seen=set()
        for item in reviews:
            if not isinstance(item,dict):fail('Revisao de P&ID malformada.')
            name=item.get('output_file')
            if name not in refs or name in seen:
                fail('Revisao P&ID duplicada ou origem sem inventario.')
            seen.add(name)
            if item.get('output_sha256')!=refs[name]['output_sha256']:
                fail('Revisao P&ID com hash divergente.')
            if item.get('review_status') not in ('REVISADO_COM_RESSALVAS','DESCARTADO'):
                fail('Status de revisao P&ID invalido; nao ha aprovacao automatica.')
            if not isinstance(item.get('reviewer'),str) or not item['reviewer'].strip():
                fail('Revisao P&ID exige responsavel humano.')
            if not isinstance(item.get('comment'),str) or not item['comment'].strip():
                fail('Revisao P&ID exige conclusao escrita.')
    for item in a['findings']:validate_finding(w,s,item)
    a.setdefault('project','NAO IDENTIFICADO')
    a.setdefault('ss','NAO IDENTIFICADA')
    for key in ('project','ss'):
        if not a[key].strip().upper().startswith(('NAO IDENTIFICAD','A DEFINIR')):
            proof=a.get('identity_evidence',{}).get(key)
            if not proof:fail('Identificacao '+key+' exige evidencia documental.')
            validate_finding(w,s,{**proof,'classification':'DOCUMENTADO'})
            if a[key] not in proof.get('verbatim_excerpt',''):
                fail('Identificacao '+key+' nao consta no trecho da fonte.')
    a['source_manifest_sha256']=s['source_manifest_sha256']
    writejson(w/'analysis.json',a)
    parts=['# RELATORIO DE ENTENDIMENTO PRELIMINAR — MCE',
      '',f"- Projeto: **{a['project']}**",f"- SS: **{a['ss']}**",
      '- Estado: **ANALISADO — AGUARDANDO DECISAO DO USUARIO**',
      '- Escopo: entendimento inicial; nao equivale a logica de seguranca aprovada.',
      '','## Entendimento','',a['summary'],'','## Constatações']
    for f in a['findings']:
        src=f.get('source_name','SEM FONTE');ln=f.get('line','?')
        parts.append(f"- **{f['classification']}** — {f['text']} (fonte: {src}, linha/pagina: {ln})")
    if derived:
        parts.extend(['','## Revisao dos resultados dos adaptadores P&ID'])
        for item in a['pid_integration_review']:
            parts.append(f"- **{item['review_status']}** — {item['output_file']}, SHA-256 {item['output_sha256']}; " 
                         f"revisor: {item['reviewer']}; observacao: {item['comment']}")
        parts.append('- Estes resultados nao sao autorizacao de logica MCE e nao substituem os P&IDs e narrativas aprovados.')
    parts.extend(['','## Lacunas'])
    parts.extend('- '+str(x) for x in a['gaps'])
    parts.extend(['','## Gate humano','',
       '**STOP**: nao modificar a MCE ate que existam modelo aprovado, escopo e SS demonstrados e autorizacao expressa.',''])
    (w/'RELATORIO_ENTENDIMENTO.md').write_text('\n'.join(parts),encoding='utf-8')
    s['status']='ANALISADO_AGUARDANDO_DECISAO';s['analysis_at']=utc()
    writejson(w/'state.json',s)
    print(json.dumps({'status':s['status'],'project':a['project'],'ss':a['ss'],
        'gaps':len(a['gaps'])},ensure_ascii=False))


def verify_contract(w, s):
    c=readjson(w/'contract.json')
    if c['template_sha256']!=s['template_sha256']:fail('Contrato nao corresponde ao modelo mestre.')
    if c.get('model_approved') is not True or s['template_classification']!='MODELO_APROVADO':
        fail('STOP: modelo Excel nao homologado no contrato e no inventario.')
    if c.get('max_causes_per_sheet')!=50:fail('Contrato alterou limite obrigatorio de 50 causas.')
    if not c.get('matrix_ranges'):fail('Declarar linhas efetivas de cada folha MCE no contrato.')
    declared=c.get('protected_rows_by_sheet')
    if not isinstance(declared,dict) or set(declared)!=set(s['workbook_inventory']):
        fail('Contrato exige faixas de linhas protegidas para TODAS as abas.')
    for sn,ranges in declared.items():
        if not isinstance(ranges,list) or not ranges:
            fail('Pelo menos o cabecalho deve estar protegido: '+sn)
        for pair in ranges:
            if not (isinstance(pair,list) and len(pair)==2 and all(isinstance(x,int) for x in pair) and 1<=pair[0]<=pair[1]):
                fail('Faixa protegida invalida em '+sn)
    if not isinstance(c.get('frozen_sheets'),list):fail('Declarar frozen_sheets como lista.')
    per_sheet={}
    for r in c['matrix_ranges']:
        sheet=r['sheet'];first,last=r['first_data_row'],r['last_data_row']
        if not isinstance(first,int) or not isinstance(last,int) or first>last or last-first+1>50:
            fail('Folha '+sheet+' excede 50 itens ou intervalo invalido.')
        if sheet not in s['workbook_inventory']:
            fail('Folha de matriz nao existe no modelo: '+sheet)
        if last>s['workbook_inventory'][sheet]['rows']:
            fail('Intervalo MCE ultrapassa area fisica existente no modelo: '+sheet)
        rows=per_sheet.setdefault(sheet,set())
        for n in range(first,last+1):
            if n in rows:fail('Intervalos de causas sobrepostos na folha '+sheet)
            rows.add(n)
        if len(rows)>50:fail('Mais de 50 itens de causa na folha '+sheet)
    return c


def run_authorize(args):
    w=Path(args.workspace);s=load_state(w)
    if s['status']!='ANALISADO_AGUARDANDO_DECISAO':fail('Analise deve ser apresentada antes da aprovacao.')
    a=readjson(w/'analysis.json'); auth=readjson(args.authorization)
    c=verify_contract(w,s)
    for f in ('project','ss'):
        if a[f].strip().upper().startswith(('NAO IDENTIFICAD','A DEFINIR')):
            fail('Projeto/SS nao demonstrado: '+f)
        if a[f]!=auth.get(f):fail('Identificacao diverge da analise: '+f)
    if auth.get('decision')!='AUTORIZO_ELABORACAO_MCE' or not auth.get('approver'):
        fail('Aprovacao expressa e identificacao do aprovador obrigatorias.')
    if auth.get('source_manifest_sha256')!=s['source_manifest_sha256']:
        fail('Aprovacao nao cobre versoes documentais correntes.')
    s['approval']={'approver':auth['approver'],'decision':auth['decision'],
                   'approved_at':utc(),'source_manifest_sha256':s['source_manifest_sha256']}
    s['status']='AUTORIZADO';writejson(w/'state.json',s)
    print(json.dumps({'status':s['status'],'approver':auth['approver']},ensure_ascii=False))


def rel_sheet_paths(zin):
    wb=ET.fromstring(zin.read('xl/workbook.xml'))
    rel=ET.fromstring(zin.read('xl/_rels/workbook.xml.rels'))
    ids={x.get('Id'):x.get('Target') for x in rel.findall('{'+PKG+'}Relationship')}
    d={}
    for s in wb.findall('m:sheets/m:sheet',NS):
        t=ids.get(s.get('{'+RNS+'}id'))
        if t is None:fail('Relacao de aba ausente.')
        p=t.lstrip('/') if t.startswith('/') else posixpath.normpath('xl/'+t)
        d[s.get('name')]=p
    return d


def is_plain_fill(styles,sid):
    xfs=styles.find(M('cellXfs'));fills=styles.find(M('fills'))
    if xfs is None or fills is None:fail('Styles XLSX incompletos.')
    if sid<0 or sid>=len(xfs):fail('Style ID invalido.')
    fill_id=int(xfs[sid].get('fillId','0'))
    if fill_id>=len(fills):fail('Fill ID invalido.')
    el=fills[fill_id]
    pat=el.find(M('patternFill'))
    if pat is None or pat.get('patternType') in (None,'none'):return True
    if pat.get('patternType')!='solid':return False
    fg=pat.find(M('fgColor'))
    if fg is None:return False
    rgb=fg.get('rgb','').upper()
    return fg.get('rgb') is not None and rgb[-6:]=='FFFFFF'


def gray_style(styles,old_sid,cache):
    if old_sid in cache:return cache[old_sid]
    xfs=styles.find(M('cellXfs'));fills=styles.find(M('fills'))
    if not is_plain_fill(styles,old_sid):
        fail('Conflito de preenchimento funcional: estilo '+str(old_sid))
    fid=cache.get('_fill_id')
    if fid is None:
        fill=ET.SubElement(fills,M('fill'))
        pat=ET.SubElement(fill,M('patternFill'),{'patternType':'solid'})
        ET.SubElement(pat,M('fgColor'),{'rgb':GRAY})
        ET.SubElement(pat,M('bgColor'),{'indexed':'64'})
        fid=len(fills)-1;fills.set('count',str(len(fills)));cache['_fill_id']=fid
    xf=copy.deepcopy(xfs[old_sid]);xf.set('fillId',str(fid));xf.set('applyFill','1')
    xfs.append(xf);xfs.set('count',str(len(xfs)));sid=len(xfs)-1;cache[old_sid]=sid
    return sid


def update_xml_sheet(src,changes,styles,styles_cache):
    root=ET.fromstring(src)
    cells={c.get('r'):c for c in root.findall('.//m:sheetData/m:row/m:c',NS)}
    for entry in changes:
        coord=entry['cell']
        c=cells.get(coord)
        if c is None:fail('Celula deve existir fisicamente no XLSX: '+coord)
        if c.find(M('f')) is not None:fail('Formula protegida: '+coord)
        sid=int(c.get('s','0'))
        c.set('s',str(gray_style(styles,sid,styles_cache)))
        # Nao tocar nos outros elementos e metadados da celula.
        for child in list(c):
            if child.tag in (M('v'),M('is')):c.remove(child)
        c.set('t','inlineStr')
        inn=ET.SubElement(c,M('is'))
        v=ET.SubElement(inn,M('t'));v.text=entry['new_text']
        if entry['new_text'] != entry['new_text'].strip():
            v.set('{http://www.w3.org/XML/1998/namespace}space','preserve')
    ET.register_namespace('',MAIN);ET.register_namespace('r',RNS)
    return ET.tostring(root,encoding='utf-8',xml_declaration=True)


def patch_workbook(master,output,changes):
    with zipfile.ZipFile(master,'r') as zi:
        paths=rel_sheet_paths(zi)
        if 'xl/styles.xml' not in zi.namelist():fail('Planilha sem styles.xml.')
        styles=ET.fromstring(zi.read('xl/styles.xml'))
        by_sheet={}
        for item in changes:
            if item['sheet'] not in paths:fail('Aba ausente: '+item['sheet'])
            by_sheet.setdefault(paths[item['sheet']],[]).append(item)
        cache={}; replacement={}
        for sheetpath,items in by_sheet.items():
            replacement[sheetpath]=update_xml_sheet(zi.read(sheetpath),items,styles,cache)
        ET.register_namespace('',MAIN)
        replacement['xl/styles.xml']=ET.tostring(styles,encoding='utf-8',xml_declaration=True)
        with zipfile.ZipFile(output,'w') as zo:
            for info in zi.infolist():
                zo.writestr(info,replacement.get(info.filename,zi.read(info.filename)))
    return set(replacement)


def validate_evidence(w,s,ch):
    if ch.get('evidence_status')!='DOCUMENTADO':fail('Apenas relacoes documentadas podem ser lancadas: '+ch['cell'])
    doc=ch.get('source_name');x=ch.get('verbatim_excerpt')
    if not x or not doc:fail('Insercao sem trecho literal e fonte: '+ch['cell'])
    f=next((z for z in s['sources'] if z['name']==doc),None)
    if not f or f['sha256']!=ch.get('source_sha256'):
        fail('Fonte fora do inventario ou hash divergente em '+ch['cell'])
    if f['extension'] in ('.txt','.md'):
        content=(w/'sources'/doc).read_text(encoding='utf-8-sig')
        if x not in content:fail('Trecho de evidência nao confere em '+ch['cell'])
    elif ch.get('manual_evidence_verified') is not True:
        fail('Evidencia binaria exige verificacao humana em '+ch['cell'])
    if not ch.get('source_location') or not ch.get('reference_revision'):
        fail('Referencia e revisao do documento obrigatorias em '+ch['cell'])


def preflight(w,s,c,changes):
    source=master_path(w); book=load_workbook(source,read_only=False)
    whitelist={sn:set(v) for sn,v in c['editable_cells'].items()}
    if not changes:fail('Nao ha alteracoes a aplicar.')
    seen=set()
    for ch in changes:
        sn,coord=ch.get('sheet'),ch.get('cell')
        if not isinstance(sn,str) or not isinstance(coord,str):fail('Alteracao sem aba e coordenada.')
        if (sn,coord) in seen:fail('Alteracao duplicada: '+sn+'!'+coord)
        seen.add((sn,coord))
        if sn not in whitelist or coord not in whitelist[sn]:fail('Celula nao autorizada: '+sn+'!'+coord)
        row=int(re.search(r'[0-9]+$',coord).group()) if re.search(r'[0-9]+$',coord) else 0
        if sn in c['frozen_sheets'] or any(lo<=row<=hi for lo,hi in c['protected_rows_by_sheet'][sn]):
            fail('Cabecalho ou aba protegida por contrato: '+sn+'!'+coord)
        if sn not in book:fail('Aba invalida: '+sn)
        if not re.fullmatch(r'[A-Z]+[1-9][0-9]*',coord):fail('Coordenada invalida: '+coord)
        cell=book[sn][coord]
        if cell.__class__.__name__=='MergedCell' or cell.data_type=='f':fail('Formula/mescla protegida: '+sn+'!'+coord)
        if not isinstance(ch.get('new_text'),str):fail('Apenas texto pode ser substituido: '+sn+'!'+coord)
        if cell.value is not None and not isinstance(cell.value,str):fail('Celula numerica nao pode ser editada.')
        if cell.value != ch.get('old_text'):fail('Valor de origem diverge em '+sn+'!'+coord)
        if ch['new_text']==ch['old_text']:fail('Sem alteracao efetiva: '+sn+'!'+coord)
        validate_evidence(w,s,ch)
    return source


def xml_changes(a,b):
    with zipfile.ZipFile(a) as z1,zipfile.ZipFile(b) as z2:
        if set(z1.namelist())!=set(z2.namelist()):fail('ZIP members alterados no XLSX.')
        return sorted(n for n in z1.namelist() if z1.read(n)!=z2.read(n))


def audit(master,derived,changes,contract):
    w1=load_workbook(master,read_only=False,keep_links=True)
    w2=load_workbook(derived,read_only=False,keep_links=True)
    bad=[]; diff=[]
    if w1.sheetnames!=w2.sheetnames:bad.append('abas alteradas')
    expected={(x['sheet'],x['cell']):x for x in changes}
    for name in w1.sheetnames:
        a=w1[name]; b=w2[name]
        for key in ('sheet_format','sheet_properties','sheet_view','page_setup','page_margins','print_options','freeze_panes','auto_filter','data_validations','conditional_formatting','protection'):
            x,y=getattr(a,key),getattr(b,key)
            if key=='auto_filter':
                same=(x.ref==y.ref and repr(x)==repr(y))
            elif key=='conditional_formatting':
                same=x._cf_rules==y._cf_rules
            else:
                same=repr(x)==repr(y)
            if not same:
                bad.append(name+': '+key+' alterado')
        if list(a.merged_cells.ranges)!=list(b.merged_cells.ranges):bad.append(name+': mesclas alteradas')
        if str(a.row_dimensions)!=str(b.row_dimensions) or str(a.column_dimensions)!=str(b.column_dimensions):
            bad.append(name+': dimensoes alteradas')
        if len(a._charts)!=len(b._charts) or len(a._images)!=len(b._images):bad.append(name+': objetos alterados')
        if a.max_row!=b.max_row or a.max_column!=b.max_column:bad.append(name+': dimensoes fisicas alteradas')
        for row in a:
            for ca in row:
                cb=b[ca.coordinate];k=(name,ca.coordinate);ex=expected.get(k)
                if ca.value!=cb.value:
                    diff.append({'sheet':name,'cell':ca.coordinate,'before':ca.value,'after':cb.value})
                    if ex is None or cb.value!=ex['new_text']:
                        bad.append('mudanca nao autorizada '+name+'!'+ca.coordinate)
                elif copy.copy(ca.fill)!=copy.copy(cb.fill):bad.append('fill sem alteracao '+name+'!'+ca.coordinate)
                for attr in ('font','border','alignment','number_format','protection','hyperlink','comment'):
                    if copy.copy(getattr(ca,attr))!=copy.copy(getattr(cb,attr)):
                        bad.append('propriedade '+attr+' alterada em '+name+'!'+ca.coordinate)
                if ca.value!=cb.value:
                    if cb.fill.patternType!='solid' or cb.fill.fgColor.rgb not in (GRAY,'00D9D9D9'):
                        bad.append('celula editada sem cinza '+name+'!'+ca.coordinate)
    if len(diff)!=len(expected):bad.append('numero de alteracoes nao confere')
    for r in contract['matrix_ranges']:
        if r['last_data_row']-r['first_data_row']+1>50:
            bad.append('mais de 50 causas na aba '+r['sheet'])
    changed_parts=xml_changes(master,derived)
    with zipfile.ZipFile(master) as z:
        paths=rel_sheet_paths(z)
    allowed_parts={'xl/styles.xml'}|{paths[x['sheet']] for x in changes}
    if not set(changed_parts).issubset(allowed_parts):bad.append('componentes OOXML indevidamente alterados')
    return {'passed':not bad,'errors':sorted(set(bad)),'differences':diff,'changed_zip_parts':changed_parts,
            'unchanged_other_zip_parts':True}


def run_apply(args):
    w=Path(args.workspace);s=load_state(w)
    if s['status']!='AUTORIZADO' or not s['approval']:fail('STOP: falta autorizacao humana explicita.')
    if s['approval']['source_manifest_sha256']!=s['source_manifest_sha256']:
        fail('Aprovacao invalida para os documentos atuais.')
    c=verify_contract(w,s);ch=readjson(args.changes)
    if not isinstance(ch,list):fail('Alteracoes devem formar lista JSON.')
    baseline=preflight(w,s,c,ch)
    output=w/'output'/'MCE_PRELIMINAR.xlsx';changed=patch_workbook(baseline,output,ch)
    result=audit(baseline,output,ch,c)
    writejson(w/'AUDITORIA_MCE.json',result)
    if not result['passed']:
        output.unlink(missing_ok=True)
        fail('STOP: auditoria de integridade reprovada: '+str(result['errors']))
    writejson(w/'ALTERACOES_MCE.json',result['differences'])
    s['status']='MCE_PRELIMINAR_PARA_REVISAO';s['output_sha256']=digest(output)
    s['artifacts'].append('output/MCE_PRELIMINAR.xlsx');writejson(w/'state.json',s)
    print(json.dumps({'status':s['status'],'cells_edited':len(ch),
         'output':str(output),'audit_pass':result['passed']},ensure_ascii=False))


def run_audit(args):
    w=Path(args.workspace);s=load_state(w)
    if s['status']!='MCE_PRELIMINAR_PARA_REVISAO':fail('MCE ainda nao foi editada.')
    p=w/'output'/'MCE_PRELIMINAR.xlsx'
    if digest(p)!=s['output_sha256']:fail('STOP: MCE editada externamente apos validacao.')
    diffs=readjson(w/'ALTERACOES_MCE.json')
    edits=[{'sheet':x['sheet'],'cell':x['cell'],'new_text':x['after']} for x in diffs]
    result=audit(master_path(w),p,edits,readjson(w/'contract.json'))
    writejson(w/'AUDITORIA_MCE.json',result)
    if not result['passed']:fail('Falha de auditoria.')
    print(json.dumps({'audit_pass':True,'changes':len(diffs)},ensure_ascii=False))


def main():
    p=argparse.ArgumentParser(description=__doc__);sub=p.add_subparsers(dest='cmd',required=True)
    a=sub.add_parser('init');a.add_argument('--workspace',required=True);a.add_argument('--master',required=True)
    a.add_argument('--sources',nargs='*',default=[]);a.add_argument('--classification',
        choices=['EXEMPLO_NAO_HOMOLOGADO','MODELO_FORNECIDO_NAO_VERIFICADO','MODELO_APROVADO'],
        default='MODELO_FORNECIDO_NAO_VERIFICADO');a.set_defaults(func=run_init)
    for name,fn,arg in [('understand',run_understand,'analysis'),('authorize',run_authorize,'authorization'),
                         ('apply',run_apply,'changes'),('audit',run_audit,None)]:
        a=sub.add_parser(name);a.add_argument('--workspace',required=True)
        if arg:a.add_argument('--'+arg,required=True)
        a.set_defaults(func=fn)
    a=sub.add_parser('status');a.add_argument('--workspace',required=True)
    a.set_defaults(func=lambda args: print(json.dumps(load_state(Path(args.workspace)),indent=2,ensure_ascii=False)))
    args=p.parse_args()
    try:args.func(args)
    except (ValueError,FileNotFoundError,KeyError,zipfile.BadZipFile) as e:
        print('BLOQUEADO: '+str(e),file=sys.stderr);sys.exit(2)


if __name__=='__main__':main()
