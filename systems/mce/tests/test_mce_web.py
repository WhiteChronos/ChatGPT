"""Private web boundary tests; synthetic data never represents engineering acceptance."""
import hashlib
import importlib.util
import json
import shutil
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]


def module():
    location = ROOT / 'scripts/mce_web.py'
    assert location.is_file(), 'MCE private web adapter is missing'
    spec = importlib.util.spec_from_file_location('mce_web_test', location)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def instance(tmp_path):
    mod = module()
    root = tmp_path / 'authorized'
    (root / 'inbox').mkdir(parents=True)
    (root / 'projects').mkdir()
    approvals = tmp_path / 'approvals'
    approvals.mkdir()
    from cryptography.hazmat.primitives.asymmetric.ed25519 import Ed25519PrivateKey
    key = Ed25519PrivateKey.generate()
    return mod, mod.MceWeb(root, approvals, key.public_key()), key


def test_status_is_independent_and_never_claims_acceptance(tmp_path):
    _, app, _ = instance(tmp_path)
    assert app.status()['product'] == 'MCE_INDEPENDENTE'
    assert app.status()['production_complete'] is False
    assert not app.status()['remote_self_approval']


@pytest.mark.parametrize('name', ['../outside.xlsx', '/tmp/a.xlsx', 'C:/x.xlsx', '.env', 'nested/../../a.xlsx'])
def test_input_escape_and_sensitive_names_are_blocked(tmp_path, name):
    _, app, _ = instance(tmp_path)
    with pytest.raises(ValueError):
        app.inspect_model(name)


def test_model_inventory_delegates_and_does_not_change_source(tmp_path):
    _, app, _ = instance(tmp_path)
    path = app.root / 'inbox/model.xlsx'
    wb = Workbook()
    wb.active['A1'] = 'SYNTHETIC MODEL'
    wb.active['B2'] = '=1+1'
    wb.save(path)
    before = path.read_bytes()
    report = app.inspect_model('model.xlsx')
    assert report['inventory']['Sheet']['formulas'] == 1
    assert report['sha256'] == hashlib.sha256(before).hexdigest()
    assert path.read_bytes() == before


def test_links_and_oversized_inputs_are_blocked(tmp_path):
    _, app, _ = instance(tmp_path)
    outside = tmp_path / 'outside.xlsx'
    outside.write_bytes(b'not a workbook')
    link = app.root / 'inbox/hard.xlsx'
    import os
    os.link(outside, link)
    with pytest.raises(ValueError):
        app.inspect_model('hard.xlsx')
    big = app.root / 'inbox/big.xlsx'
    with big.open('wb') as handle:
        handle.truncate(21 * 1024 * 1024)
    with pytest.raises(ValueError):
        app.inspect_model('big.xlsx')


def test_zip_xml_entities_and_sparse_dimension_are_blocked(tmp_path):
    _, app, _ = instance(tmp_path)
    import zipfile
    for content in [b'<!DOCTYPE x [<!ENTITY e "bad">]><worksheet>&e;</worksheet>',
                    b'<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><dimension ref="A1:XFD1048576"/></worksheet>']:
        path = app.root / 'inbox/evil.xlsx'
        with zipfile.ZipFile(path, 'w') as z:
            z.writestr('xl/worksheets/sheet1.xml', content)
        with pytest.raises(ValueError):
            app.inspect_model('evil.xlsx')


def prepared(tmp_path):
    mod, app, key = instance(tmp_path)
    spec = importlib.util.spec_from_file_location('existing_mce_fixture', ROOT / 'tests/test_mce_workflow.py')
    fixtures = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(fixtures)
    source = tmp_path / 'fixture'
    source.mkdir()
    ws, master, _, _, changes, _ = fixtures.fixture(source)
    shutil.copytree(ws, app.root / 'projects/demo')
    shutil.copy2(changes, app.root / 'inbox/changes.json')
    return mod, app, key, master


def approve(mod, app, key):
    request = app.approval_request('demo', 'changes.json')
    envelope = mod.sign_approval(request, key, 'TEST OPERATOR', lifetime_seconds=300)
    (app.approvals / 'approved.json').write_text(json.dumps(envelope), encoding='utf-8')
    return envelope


def test_cannot_apply_without_external_signature(tmp_path):
    _, app, _, master = prepared(tmp_path)
    before = master.read_bytes()
    with pytest.raises(ValueError):
        app.apply('demo', 'changes.json', 'missing.json')
    assert master.read_bytes() == before
    assert not (app.root / 'projects/demo/output/MCE_PRELIMINAR.xlsx').exists()


def test_signed_exact_change_generates_copy_and_preserves_master(tmp_path):
    mod, app, key, master = prepared(tmp_path)
    before = master.read_bytes()
    approve(mod, app, key)
    result = app.apply('demo', 'changes.json', 'approved.json')
    assert result['status'] == 'MCE_PRELIMINAR_PARA_REVISAO'
    assert result['audit_pass'] is True
    assert result['production_complete'] is False
    output = app.root / result['output']
    wb = load_workbook(output)
    assert wb['CAUSAS']['C2'].value == 'Deteccao de H2 confirmada'
    assert wb['CAUSAS']['D2'].value == '=1+1'
    assert wb['CAUSAS']['C2'].fill.fgColor.rgb == 'FFD9D9D9'
    assert master.read_bytes() == before
    with pytest.raises(ValueError):
        app.apply('demo', 'changes.json', 'approved.json')


@pytest.mark.parametrize('tamper', ['changes', 'contract', 'signature', 'analysis'])
def test_approval_cannot_survive_tampering(tmp_path, tamper):
    mod, app, key, _ = prepared(tmp_path)
    envelope = approve(mod, app, key)
    if tamper == 'signature':
        envelope['signature'] = '00' * 64
        (app.approvals / 'approved.json').write_text(json.dumps(envelope), encoding='utf-8')
    else:
        path = app.root / ('inbox/changes.json' if tamper == 'changes' else f'projects/demo/{tamper}.json')
        path.write_text(path.read_text(encoding='utf-8') + ' ', encoding='utf-8')
    with pytest.raises(ValueError):
        app.apply('demo', 'changes.json', 'approved.json')
    assert not (app.root / 'projects/demo/output/MCE_PRELIMINAR.xlsx').exists()


def test_expired_and_wrong_operator_approvals_are_rejected(tmp_path):
    mod, app, key, _ = prepared(tmp_path)
    import datetime as dt
    from cryptography.hazmat.primitives.asymmetric.ed25519 import Ed25519PrivateKey
    request = app.approval_request('demo', 'changes.json')
    envelope = mod.sign_approval(request, Ed25519PrivateKey.generate(), 'UNTRUSTED')
    path = app.approvals / 'approved.json'
    path.write_text(json.dumps(envelope), encoding='utf-8')
    with pytest.raises(ValueError):
        app.apply('demo', 'changes.json', 'approved.json')
    envelope = mod.sign_approval(request, key, 'EXPIRED')
    envelope['payload']['issued_at'] = (dt.datetime.now(dt.timezone.utc) - dt.timedelta(hours=2)).isoformat()
    envelope['payload']['expires_at'] = (dt.datetime.now(dt.timezone.utc) - dt.timedelta(hours=1)).isoformat()
    envelope['signature'] = key.sign(mod.canonical(envelope['payload'])).hex()
    path.write_text(json.dumps(envelope), encoding='utf-8')
    with pytest.raises(ValueError):
        app.apply('demo', 'changes.json', 'approved.json')


def test_workspace_cannot_redirect_canonical_source_reads(tmp_path):
    _, app, _, _ = prepared(tmp_path)
    path = app.root / 'projects/demo/state.json'
    value = json.loads(path.read_text(encoding='utf-8'))
    value['sources'][0]['name'] = '../../../outside.txt'
    path.write_text(json.dumps(value), encoding='utf-8')
    with pytest.raises(ValueError):
        app.approval_request('demo', 'changes.json')


def test_approval_directory_cannot_be_inside_remote_write_root(tmp_path):
    mod, app, _ = instance(tmp_path)
    inside = app.root / 'approvals'
    inside.mkdir()
    with pytest.raises(ValueError):
        mod.MceWeb(app.root, inside, None)


def test_initialization_and_understanding_use_canonical_stop_gate(tmp_path):
    mod, app, key, _ = prepared(tmp_path)
    old = app.root / 'projects/demo'
    shutil.copyfile(old / 'base/modelo.xlsx', app.root / 'inbox/model.xlsx')
    shutil.copyfile(old / 'sources/fonte.txt', app.root / 'inbox/fonte.txt')
    shutil.copyfile(old / 'analysis.json', app.root / 'inbox/analysis.json')
    shutil.copyfile(old / 'contract.json', app.root / 'inbox/contract.json')
    report = app.initialize('new-project', 'model.xlsx', ['fonte.txt'])
    assert report['status'] == 'INVENTARIADO'
    report = app.understand('new-project', 'analysis.json', 'contract.json')
    assert report['status'] == 'ANALISADO_AGUARDANDO_DECISAO'
    assert not (app.root / 'projects/new-project/output/MCE_PRELIMINAR.xlsx').exists()
    assert app.approval_request('new-project', 'changes.json')['project_id'] == 'new-project'
    request = app.approval_request('new-project', 'changes.json')
    envelope = mod.sign_approval(request, key, 'SYNTHETIC MODEL REVIEWER')
    (app.approvals / 'new.json').write_text(json.dumps(envelope), encoding='utf-8')
    result = app.apply('new-project', 'changes.json', 'new.json')
    assert result['audit_pass']
    original_state = json.loads((app.root / 'projects/new-project/state.json').read_text(encoding='utf-8'))
    assert original_state['template_classification'] == 'MODELO_FORNECIDO_NAO_VERIFICADO'


def test_total_workbook_dimension_budget_is_bounded(tmp_path):
    mod, app, _ = instance(tmp_path)
    import zipfile
    path = app.root / 'inbox/huge-dimension.xlsx'
    with zipfile.ZipFile(path, 'w') as zipped:
        zipped.writestr('xl/worksheets/sheet1.xml', '<worksheet><dimension ref="A1:IV20000"/></worksheet>')
    with pytest.raises(ValueError):
        mod.inspect_zip(path)


def test_operator_cli_signs_only_outside_remote_root_and_never_overwrites(tmp_path):
    mod, app, key, _ = prepared(tmp_path)
    import subprocess
    import sys
    from cryptography.hazmat.primitives.serialization import Encoding, PrivateFormat, NoEncryption
    request = app.approval_request('demo', 'changes.json')
    request_path = tmp_path / 'reviewed.json'
    request_path.write_bytes(mod.canonical(request))
    key_path = tmp_path / 'operator.pem'
    private_data = key.private_bytes(Encoding.PEM, PrivateFormat.PKCS8, NoEncryption())
    key_path.write_bytes(private_data)
    output = app.approvals / 'cli.json'
    args = [sys.executable, str(ROOT / 'scripts/mce_operator_approval.py'),
            '--request', str(request_path), '--private-key', str(key_path),
            '--authorized-root', str(app.root), '--output', str(output),
            '--approver', 'SYNTHETIC OPERATOR', '--decision', 'AUTORIZO_ELABORACAO_MCE', '--model-reviewed']
    result = subprocess.run(args, capture_output=True)
    assert result.returncode == 0
    envelope = json.loads(output.read_bytes())
    key.public_key().verify(bytes.fromhex(envelope['signature']), mod.canonical(envelope['payload']))
    assert private_data not in result.stdout + result.stderr
    before = output.read_bytes()
    assert subprocess.run(args, capture_output=True).returncode != 0
    assert output.read_bytes() == before
    args[args.index('--output') + 1] = str(app.root / 'forbidden.json')
    assert subprocess.run(args, capture_output=True).returncode != 0
    assert not (app.root / 'forbidden.json').exists()


def test_web_runtime_keeps_windows_only_contract(monkeypatch):
    mod = module()
    monkeypatch.setattr(mod.sys, 'platform', 'linux')
    with pytest.raises(ValueError, match='WINDOWS_ONLY'):
        mod.require_windows()
