import importlib.util
import json
import contextlib
import io
from types import SimpleNamespace
import sys
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill

ROOT = Path(__file__).resolve().parents[1]
CLI = ROOT/'.agents/skills/engenharia-mce-projeto/scripts/mce_workflow.py'
SPEC = importlib.util.spec_from_file_location('mce_workflow',CLI)
mod = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(mod)


def command(*args, ok=True):
    saved=sys.argv
    out,err=io.StringIO(),io.StringIO()
    code=0
    try:
        sys.argv=[str(CLI),*map(str,args)]
        with contextlib.redirect_stdout(out), contextlib.redirect_stderr(err):
            try:
                mod.main()
            except SystemExit as e:
                code=int(e.code or 0)
    finally:
        sys.argv=saved
    result=SimpleNamespace(returncode=code,stdout=out.getvalue(),stderr=err.getvalue())
    if ok: assert result.returncode==0,(result.stdout,result.stderr)
    else: assert result.returncode!=0,(result.stdout,result.stderr)
    return result


def dump(path,obj):
    path.write_text(json.dumps(obj,ensure_ascii=False),encoding='utf-8')
    return path


def fixture(tmp_path,colored=False):
    wb=Workbook();a=wb.active;a.title='CONTROLE';a['A1']='CABECALHO INALTERAVEL'
    b=wb.create_sheet('CAUSAS');b['C2']='ANTIGO';b['D2']='=1+1'
    if colored:b['C2'].fill=PatternFill('solid',fgColor='FF0000')
    c=wb.create_sheet('MCE');c['A4']='ID_CAUSA';c['A5']='C01';c['E5']='?';c['A6']='C02'
    model=tmp_path/'modelo.xlsx';wb.save(model)
    source=tmp_path/'fonte.txt';source.write_text('Projeto PROJETO TESTE e SS-TESTE\nDeteccao H2 documentada\n',encoding='utf-8')
    ws=tmp_path/'projeto';command('init','--workspace',ws,'--master',model,'--sources',source,
          '--classification','MODELO_APROVADO')
    s=mod.load_state(ws)
    analysis={'project':'PROJETO TESTE','ss':'SS-TESTE','summary':'Sistema de teste sintético.',
         'identity_evidence':{
            'project':{'source_name':source.name,'line':1,'verbatim_excerpt':'Projeto PROJETO TESTE e SS-TESTE'},
            'ss':{'source_name':source.name,'line':1,'verbatim_excerpt':'Projeto PROJETO TESTE e SS-TESTE'}},
         'findings':[{'classification':'DOCUMENTADO','text':'Deteccao H2 esta documentada',
             'source_name':source.name,'line':2,'verbatim_excerpt':'Deteccao H2 documentada'}],
         'gaps':['Setpoint nao documentado']}
    command('understand','--workspace',ws,'--analysis',dump(tmp_path/'analise.json',analysis))
    contract=mod.readjson(ws/'contract.json')
    contract.update(model_approved=True,
        matrix_ranges=[{'sheet':'MCE','first_data_row':5,'last_data_row':6}],
        editable_cells={'CAUSAS':['C2'],'MCE':['E5']},
        protected_rows_by_sheet={'CONTROLE':[[1,1]],'CAUSAS':[[1,1]],'MCE':[[1,4]]},
        frozen_sheets=['CONTROLE'])
    mod.writejson(ws/'contract.json',contract)
    approval={'decision':'AUTORIZO_ELABORACAO_MCE','approver':'Aprovador de teste',
              'project':'PROJETO TESTE','ss':'SS-TESTE',
              'source_manifest_sha256':s['source_manifest_sha256']}
    ap=dump(tmp_path/'aprovacao.json',approval)
    ch=[{'sheet':'CAUSAS','cell':'C2','old_text':'ANTIGO','new_text':'Deteccao de H2 confirmada',
        'source_name':source.name,'source_sha256':s['sources'][0]['sha256'],
        'source_location':'linha 2','reference_revision':'R00',
        'verbatim_excerpt':'Deteccao H2 documentada','evidence_status':'DOCUMENTADO'}]
    changefile=dump(tmp_path/'mudancas.json',ch)
    return ws,model,source,ap,changefile,ch


def test_end_to_end_preserves_unedited_ooxml(tmp_path):
    ws,model,source,ap,changes,ch=fixture(tmp_path)
    command('apply','--workspace',ws,'--changes',changes,ok=False)  # sem gate
    command('authorize','--workspace',ws,'--authorization',ap)
    command('apply','--workspace',ws,'--changes',changes)
    command('audit','--workspace',ws)
    base=ws/'base/modelo.xlsx';out=ws/'output/MCE_PRELIMINAR.xlsx'
    assert mod.digest(model)==mod.digest(base)
    w=load_workbook(out)
    assert w['CAUSAS']['C2'].value=='Deteccao de H2 confirmada'
    assert w['CAUSAS']['C2'].fill.fgColor.rgb=='FFD9D9D9'
    assert w['CAUSAS']['D2'].value=='=1+1'
    assert w['CONTROLE']['A1'].value=='CABECALHO INALTERAVEL'
    parts=mod.xml_changes(base,out)
    assert set(parts)=={'xl/styles.xml','xl/worksheets/sheet2.xml'}
    assert mod.readjson(ws/'AUDITORIA_MCE.json')['passed']


def test_blocks_unauthorized_cell_and_formula(tmp_path):
    ws,_,_,ap,changes,ch=fixture(tmp_path)
    command('authorize','--workspace',ws,'--authorization',ap)
    ch[0].update(cell='A1',old_text=None)
    dump(changes,ch)
    r=command('apply','--workspace',ws,'--changes',changes,ok=False)
    assert 'nao autorizada' in r.stderr
    ch[0].update(cell='D2',old_text='=1+1')
    mod.writejson(ws/'contract.json',{**mod.readjson(ws/'contract.json'),
          'editable_cells':{'CAUSAS':['D2']}})
    dump(changes,ch)
    r=command('apply','--workspace',ws,'--changes',changes,ok=False)
    assert 'Formula' in r.stderr


def test_blocks_51_items(tmp_path):
    ws,_,_,ap,_,_=fixture(tmp_path)
    c=mod.readjson(ws/'contract.json');c['matrix_ranges'][0]['last_data_row']=55
    mod.writejson(ws/'contract.json',c)
    r=command('authorize','--workspace',ws,'--authorization',ap,ok=False)
    assert '50 itens' in r.stderr


def test_blocks_colored_cell(tmp_path):
    ws,_,_,ap,changes,_=fixture(tmp_path,colored=True)
    command('authorize','--workspace',ws,'--authorization',ap)
    r=command('apply','--workspace',ws,'--changes',changes,ok=False)
    assert 'preenchimento funcional' in r.stderr
    assert not (ws/'output/MCE_PRELIMINAR.xlsx').exists()


def test_blocks_unverified_source_and_unapproved_model(tmp_path):
    ws,_,source,ap,changes,_=fixture(tmp_path)
    c=mod.readjson(ws/'contract.json');c['model_approved']=False
    mod.writejson(ws/'contract.json',c)
    r=command('authorize','--workspace',ws,'--authorization',ap,ok=False)
    assert 'nao homologado' in r.stderr
    source_copy=ws/'sources'/source.name
    source_copy.write_text('ALTERADO',encoding='utf-8')
    r=command('status','--workspace',ws,ok=False)
    assert 'fonte alterada' in r.stderr


def test_audit_rejects_header_mutation(tmp_path):
    ws,_,_,ap,changes,ch=fixture(tmp_path)
    command('authorize','--workspace',ws,'--authorization',ap)
    command('apply','--workspace',ws,'--changes',changes)
    output=ws/'output/MCE_PRELIMINAR.xlsx'
    w=load_workbook(output);w['CONTROLE']['A1']='CORROMPIDO';w.save(output)
    report=mod.audit(ws/'base/modelo.xlsx',output,ch,mod.readjson(ws/'contract.json'))
    assert not report['passed']
    assert any('nao autorizada' in e or 'componentes OOXML' in e for e in report['errors'])


def test_header_blocked_even_when_whitelisted(tmp_path):
    ws,_,source,ap,changes,ch=fixture(tmp_path)
    command('authorize','--workspace',ws,'--authorization',ap)
    c=mod.readjson(ws/'contract.json');c['editable_cells']['CONTROLE']=['A1']
    mod.writejson(ws/'contract.json',c)
    ch[0].update(sheet='CONTROLE',cell='A1',old_text='CABECALHO INALTERAVEL')
    dump(changes,ch)
    r=command('apply','--workspace',ws,'--changes',changes,ok=False)
    assert 'Cabecalho ou aba protegida' in r.stderr


def test_analysis_requires_documented_project_and_ss(tmp_path):
    model=tmp_path/'modelo.xlsx';a=Workbook();a.save(model)
    doc=tmp_path/'fonte.txt';doc.write_text('Sistema sem identificacao formal.',encoding='utf-8')
    ws=tmp_path/'proj';command('init','--workspace',ws,'--master',model,'--sources',doc)
    analysis={'project':'PROJETO INVENTADO','ss':'SS-INVENTADA','summary':'Resumo sintético',
              'findings':[],'gaps':['pendencia']}
    r=command('understand','--workspace',ws,'--analysis',dump(tmp_path/'analysis.json',analysis),ok=False)
    assert 'exige evidencia documental' in r.stderr


def test_unverified_evidence_rejected_at_apply(tmp_path):
    ws,_,_,ap,changes,ch=fixture(tmp_path)
    command('authorize','--workspace',ws,'--authorization',ap)
    ch[0]['verbatim_excerpt']='TRECHO INVENTADO'
    dump(changes,ch)
    r=command('apply','--workspace',ws,'--changes',changes,ok=False)
    assert 'nao confere' in r.stderr
