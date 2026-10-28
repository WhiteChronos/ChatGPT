#!/usr/bin/env python3
"""Inventario somente-leitura da MCE mestre: SHA-256, abas, formulas e limites."""
import sys, json, hashlib
from pathlib import Path
from openpyxl import load_workbook

if len(sys.argv)!=2:
    sys.exit('Uso: inventariar_modelo.py modelo.xlsx')
p=Path(sys.argv[1]);b=p.read_bytes();w=load_workbook(p,read_only=False,keep_links=True)
r={'file':p.name,'sha256':hashlib.sha256(b).hexdigest(),'bytes':len(b),'sheets':[]}
for s in w:
    formulas=sum(c.data_type=='f' for row in s for c in row)
    r['sheets'].append({'name':s.title,'max_row':s.max_row,'max_col':s.max_column,
        'formulas':formulas,'merged':sorted(map(str,s.merged_cells.ranges)),
        'protected':bool(s.protection.sheet),'charts':len(s._charts),'images':len(s._images),
        'print_area':str(s.print_area or '')})
print(json.dumps(r,indent=2,ensure_ascii=False))
