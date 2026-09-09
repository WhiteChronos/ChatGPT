"""Gera o Excel de controle no padrão aprovado.

Regras:
- dúvidas/perguntas nunca entram no Excel;
- todos os objetivos formais e erros confirmados começam em ☐;
- coluna de status centralizada em 22 pt;
- borda externa preta grossa e internas pretas finas;
- texto com quebra automática e altura confortável;
- sem colunas de responsável/evidência de atendimento/data;
- sem abas Confronto_IO ou Verificacao.
"""

from __future__ import annotations

import json
from copy import copy
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

from pipeline.comment_control_pipeline import gate

HEADERS = [
    "ID",
    "Grau",
    "Documento(s)",
    "Objetivo / Erro",
    "O que precisa ser verificado / atendido",
    "Evidência / constatação na revisão analisada",
    "Fonte / localização",
    "Comentário atendido",
]

THIN = Side(style="thin", color="000000")
THICK = Side(style="thick", color="000000")
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
BODY_FONT = Font(color="000000", size=11)
CHECK_FONT = Font(color="000000", size=22, bold=True)
ERROR_FILL = PatternFill("solid", fgColor="F4CCCC")
ERROR_FONT = Font(color="9C0006", bold=True, size=11)
SECTION_FILL = PatternFill("solid", fgColor="D9EAF7")


def document_location(record: dict[str, Any]) -> str:
    parts = [record.get("document_code"), record.get("revision"), record.get("page_or_item")]
    return " — ".join(str(p).strip() for p in parts if p)


def visual_status(record: dict[str, Any]) -> str:
    return "☑" if record.get("status_control") == "CHECKED" else "☐"


def _apply_table_border(ws, min_row: int, max_row: int, min_col: int, max_col: int) -> None:
    for row in ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
        for cell in row:
            top = THICK if cell.row == min_row else THIN
            bottom = THICK if cell.row == max_row else THIN
            left = THICK if cell.column == min_col else THIN
            right = THICK if cell.column == max_col else THIN
            cell.border = Border(top=top, bottom=bottom, left=left, right=right)


def _row_height(values: list[str]) -> float:
    longest = max((len(value or "") for value in values), default=0)
    approx_lines = max(2, min(10, (longest // 75) + 1))
    return max(36.0, approx_lines * 18.0)


def _records(payload: dict[str, Any]) -> list[dict[str, Any]]:
    return [*payload.get("comments", []), *payload.get("new_divergences", [])]


def _setup_main_sheet(ws, payload: dict[str, Any]) -> None:
    ws.title = "Controle_Geral"
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A5"

    formal_count = len(payload.get("comments", []))
    error_count = len(payload.get("new_divergences", []))
    ws.merge_cells("A1:H1")
    ws["A1"] = "CONTROLE DE VERIFICAÇÃO"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    ws.merge_cells("A2:H2")
    ws["A2"] = (
        f"{formal_count} objetivos formais + {error_count} erros confirmados. "
        "Dúvidas são resolvidas antes da emissão e nunca aparecem nesta planilha. Todos iniciam em ☐."
    )
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="center")
    ws["A2"].font = Font(size=10, italic=True)
    ws.row_dimensions[2].height = 34

    for idx, value in enumerate(HEADERS, start=1):
        cell = ws.cell(4, idx, value)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[4].height = 42

    status_validation = DataValidation(type="list", formula1='"☐,☑"', allow_blank=False)
    status_validation.error = "Selecione somente ☐ ou ☑."
    status_validation.errorTitle = "Status inválido"
    status_validation.prompt = "☐ = não confirmado | ☑ = confirmado"
    status_validation.promptTitle = "Comentário atendido"
    ws.add_data_validation(status_validation)

    for record in _records(payload):
        location = document_location(record)
        objective = str(record.get("original_comment") or "")
        action = str(record.get("compiled_action") or "")
        finding_basis = str(record.get("finding_basis") or "")
        source_location = str(record.get("source_location") or "")
        ws.append([
            record.get("comment_id"),
            record.get("severity"),
            location,
            objective,
            action,
            finding_basis,
            source_location,
            visual_status(record),
        ])
        r = ws.max_row
        ws.row_dimensions[r].height = _row_height([location, objective, action, finding_basis, source_location])

        for c in range(1, 8):
            ws.cell(r, c).font = BODY_FONT
            ws.cell(r, c).alignment = Alignment(vertical="top", wrap_text=True)

        if record.get("origin_type") == "NEW_DIVERGENCE":
            ws.cell(r, 4).fill = ERROR_FILL
            ws.cell(r, 4).font = ERROR_FONT

        status_cell = ws.cell(r, 8)
        status_cell.font = CHECK_FONT
        status_cell.alignment = Alignment(horizontal="center", vertical="center")
        status_validation.add(status_cell)

    ws.column_dimensions["A"].width = 13
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 36
    ws.column_dimensions["D"].width = 55
    ws.column_dimensions["E"].width = 72
    ws.column_dimensions["F"].width = 68
    ws.column_dimensions["G"].width = 45
    ws.column_dimensions["H"].width = 22

    if ws.max_row >= 4:
        _apply_table_border(ws, 4, ws.max_row, 1, 8)

    ws.auto_filter.ref = f"A4:H{ws.max_row}"
    ws.print_title_rows = "1:4"
    ws.print_area = f"A1:H{ws.max_row}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def _setup_lists(wb: Workbook) -> None:
    ws = wb.create_sheet("Listas")
    ws["A1"] = "☐"
    ws["A2"] = "☑"
    ws.column_dimensions["A"].width = 10


def _setup_rules(wb: Workbook) -> None:
    ws = wb.create_sheet("Fontes_e_Regras")
    ws.sheet_view.showGridLines = False
    rows = [
        ["REGRA", "DESCRIÇÃO"],
        ["Pré-verificação", "Toda dúvida deve ser sanada antes de gerar Excel, Word, PDF ou Power BI."],
        ["Perguntas", "Perguntas e itens DÚVIDA/A CONFIRMAR são proibidos na planilha final."],
        ["Erro", "Somente não conformidade diretamente observável ou demonstrável é apresentada como erro."],
        ["Quantidade", "100% dos comentários formais da origem devem permanecer no controle."],
        ["Status", "Todos os itens iniciam em ☐; ☑ somente após confirmação humana."],
        ["Impressão", "Espaço em branco por menor quantidade de conteúdo não é erro por si só."],
        ["Documentos diferentes", "ET, FD, LI, MD e DE não precisam repetir todo o conteúdo uns dos outros."],
        ["Similar técnico", "Referências comerciais diferentes com 'ou similar técnico' não são erro por si só."],
    ]
    for row in rows:
        ws.append(row)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.font = BODY_FONT
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 100
    _apply_table_border(ws, 1, ws.max_row, 1, 2)


def build_workbook(payload: dict[str, Any]) -> Workbook:
    # gate também impede geração quando existir clarification_question aberta.
    gate(payload)
    wb = Workbook()
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"

    _setup_main_sheet(wb.active, payload)
    _setup_lists(wb)
    _setup_rules(wb)
    return wb


def export_excel(payload: dict[str, Any], output_path: str | Path) -> Path:
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb = build_workbook(payload)
    wb.save(output_path)
    return output_path


if __name__ == "__main__":
    import argparse

    parser = argparse.ArgumentParser()
    parser.add_argument("input")
    parser.add_argument("output")
    args = parser.parse_args()

    payload = json.loads(Path(args.input).read_text(encoding="utf-8"))
    print(export_excel(payload, args.output))
