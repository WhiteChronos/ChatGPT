#!/usr/bin/env python3
"""Controle fail-closed de revisão de LI - Lista de Material.

O script não calcula quantitativos de engenharia. Ele protege histórico, estrutura e
rastreabilidade do workbook. A edição de valores da nova revisão continua dependente
das fontes autorizadas pelo projeto.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import re
import shutil
import zipfile
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

REV_RE = re.compile(r"^REV\.\s*(.+)$", re.IGNORECASE)


class LIControlError(RuntimeError):
    pass


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for block in iter(lambda: f.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def normalize_revision(value: Any) -> str | None:
    if value is None:
        return None
    text = " ".join(str(value).strip().upper().split())
    m = REV_RE.match(text)
    if not m:
        return None
    token = m.group(1).strip()
    return f"REV. {token}"


def workbook_is_valid_ooxml(path: Path) -> bool:
    if not zipfile.is_zipfile(path):
        return False
    with zipfile.ZipFile(path) as zf:
        names = set(zf.namelist())
    return "[Content_Types].xml" in names and "xl/workbook.xml" in names


def _column_widths(ws) -> dict[str, float | None]:
    """Captura larguras explicitamente configuradas, inclusive colunas futuras vazias.

    ws.max_column cresce quando a nova revisão recebe conteúdo. Usá-lo como limite faria uma
    coluna já formatada, porém vazia na baseline, parecer uma alteração estrutural quando a revisão
    seguinte fosse preenchida. A estrutura protegida é o conjunto de ColumnDimension persistido no
    workbook, não a área atualmente ocupada por valores.
    """
    result: dict[str, float | None] = {}
    for key, dim in ws.column_dimensions.items():
        if dim.width is not None:
            result[str(key)] = float(dim.width)
    return result


def _row_heights(ws) -> dict[str, float | None]:
    result: dict[str, float | None] = {}
    for idx, dim in ws.row_dimensions.items():
        if dim.height is not None:
            result[str(idx)] = float(dim.height)
    return result


def find_revision_headers(ws, scan_rows: int = 40) -> list[dict[str, Any]]:
    headers: list[dict[str, Any]] = []
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, scan_rows)):
        for cell in row:
            label = normalize_revision(cell.value)
            if label:
                headers.append({"label": label, "cell": cell.coordinate, "row": cell.row, "col": cell.column})
    return sorted(headers, key=lambda x: (x["row"], x["col"]))


def capture_revision_column(ws, header: dict[str, Any]) -> dict[str, Any]:
    col = int(header["col"])
    row0 = int(header["row"]) + 1
    values: dict[str, Any] = {}
    for row in range(row0, ws.max_row + 1):
        cell = ws.cell(row=row, column=col)
        values[cell.coordinate] = cell.value
    return {
        "label": header["label"],
        "header_cell": header["cell"],
        "row": header["row"],
        "col": header["col"],
        "values": values,
    }


def snapshot_workbook(path: Path) -> dict[str, Any]:
    if not workbook_is_valid_ooxml(path):
        raise LIControlError(f"XLSX/OOXML inválido: {path}")
    wb = load_workbook(path, data_only=False, keep_links=True)
    sheets: dict[str, Any] = {}
    for ws in wb.worksheets:
        headers = find_revision_headers(ws)
        sheets[ws.title] = {
            "max_row": ws.max_row,
            "max_column": ws.max_column,
            "merged_ranges": sorted(str(rng) for rng in ws.merged_cells.ranges),
            "print_area": str(ws.print_area) if ws.print_area else None,
            "freeze_panes": str(ws.freeze_panes) if ws.freeze_panes else None,
            "column_widths": _column_widths(ws),
            "row_heights": _row_heights(ws),
            "revision_headers": headers,
            "revision_columns": [capture_revision_column(ws, h) for h in headers],
        }
    return {
        "file": path.name,
        "sha256": sha256_file(path),
        "sheet_order": wb.sheetnames,
        "external_links_count": len(getattr(wb, "_external_links", []) or []),
        "sheets": sheets,
    }


def _sheet_structure(snapshot: dict[str, Any], sheet: str) -> dict[str, Any]:
    data = snapshot["sheets"][sheet]
    return {
        "merged_ranges": data["merged_ranges"],
        "print_area": data["print_area"],
        "freeze_panes": data["freeze_panes"],
        "column_widths": data["column_widths"],
        "row_heights": data["row_heights"],
    }


def _revision_map(snapshot: dict[str, Any], sheet: str) -> dict[str, dict[str, Any]]:
    return {x["label"]: x for x in snapshot["sheets"][sheet]["revision_columns"]}


def compare_structure(baseline: dict[str, Any], current: dict[str, Any]) -> list[str]:
    errors: list[str] = []
    if baseline["sheet_order"] != current["sheet_order"]:
        errors.append("sheet_order_changed")
    for sheet in baseline["sheet_order"]:
        if sheet not in current["sheets"]:
            errors.append(f"missing_sheet:{sheet}")
            continue
        if _sheet_structure(baseline, sheet) != _sheet_structure(current, sheet):
            errors.append(f"sheet_structure_changed:{sheet}")
    if current["external_links_count"] > baseline["external_links_count"]:
        errors.append("unexpected_external_links_added")
    return errors


def compare_historical_revisions(
    baseline: dict[str, Any], current: dict[str, Any], target_revision: str
) -> list[str]:
    errors: list[str] = []
    target_revision = normalize_revision(target_revision) or target_revision
    for sheet in baseline["sheet_order"]:
        if sheet not in current["sheets"]:
            continue
        bmap = _revision_map(baseline, sheet)
        cmap = _revision_map(current, sheet)
        for label, old in bmap.items():
            if label == target_revision:
                continue
            new = cmap.get(label)
            if new is None:
                errors.append(f"historical_revision_missing:{sheet}:{label}")
                continue
            if old["header_cell"] != new["header_cell"]:
                errors.append(f"historical_revision_moved:{sheet}:{label}")
            if old["values"] != new["values"]:
                errors.append(f"historical_revision_changed:{sheet}:{label}")
    return errors


def validate_target_adjacency(current: dict[str, Any], previous_revision: str, target_revision: str) -> list[str]:
    errors: list[str] = []
    previous_revision = normalize_revision(previous_revision) or previous_revision
    target_revision = normalize_revision(target_revision) or target_revision
    found_pair = False
    for sheet, data in current["sheets"].items():
        headers = data["revision_headers"]
        by_row: dict[int, list[dict[str, Any]]] = {}
        for h in headers:
            by_row.setdefault(int(h["row"]), []).append(h)
        for row_headers in by_row.values():
            p = next((h for h in row_headers if h["label"] == previous_revision), None)
            t = next((h for h in row_headers if h["label"] == target_revision), None)
            if p and t:
                found_pair = True
                if int(t["col"]) != int(p["col"]) + 1:
                    errors.append(f"target_revision_not_adjacent:{sheet}:{previous_revision}->{target_revision}")
    if not found_pair:
        errors.append(f"revision_pair_not_found:{previous_revision}->{target_revision}")
    return errors


def optional_tool_coverage() -> dict[str, str]:
    tools = {
        "workbooklens": shutil.which("workbooklens") or "UNAVAILABLE",
        "sheetparity": shutil.which("sheetparity") or "UNAVAILABLE",
        "libreoffice": shutil.which("libreoffice") or shutil.which("soffice") or "UNAVAILABLE",
    }
    return tools


def validate(
    workbook: Path,
    baseline_workbook: Path | None,
    previous_revision: str,
    target_revision: str,
) -> dict[str, Any]:
    current = snapshot_workbook(workbook)
    findings: list[str] = []
    coverage = optional_tool_coverage()
    baseline = None
    if baseline_workbook:
        baseline = snapshot_workbook(baseline_workbook)
        findings.extend(compare_structure(baseline, current))
        findings.extend(compare_historical_revisions(baseline, current, target_revision))
    else:
        findings.append("baseline_missing:historical_preservation_not_proven")
    findings.extend(validate_target_adjacency(current, previous_revision, target_revision))

    reproved = [x for x in findings if x.startswith((
        "sheet_order_changed",
        "missing_sheet:",
        "sheet_structure_changed:",
        "unexpected_external_links_added",
        "historical_revision_missing:",
        "historical_revision_moved:",
        "historical_revision_changed:",
        "target_revision_not_adjacent:",
        "revision_pair_not_found:",
    ))]
    hold = [x for x in findings if x.startswith("baseline_missing:")]
    status = "REPROVADO" if reproved else ("HOLD" if hold else "VALIDADO")
    return {
        "control_id": "LI-MATERIAL-CONTROL-V1",
        "status": status,
        "workbook": current["file"],
        "workbook_sha256": current["sha256"],
        "baseline": baseline["file"] if baseline else None,
        "previous_revision": normalize_revision(previous_revision) or previous_revision,
        "target_revision": normalize_revision(target_revision) or target_revision,
        "findings": findings,
        "open_source_qa_coverage": coverage,
        "rules": ["GR-050", "GR-051", "GR-052", "GR-053", "GR-054", "GR-055", "GR-056", "GR-057", "GR-058"],
    }


def main() -> int:
    parser = argparse.ArgumentParser(description="Controle de revisão de LI - Lista de Material")
    sub = parser.add_subparsers(dest="command", required=True)

    p_snapshot = sub.add_parser("snapshot")
    p_snapshot.add_argument("workbook")
    p_snapshot.add_argument("--output", required=True)

    p_validate = sub.add_parser("validate")
    p_validate.add_argument("workbook")
    p_validate.add_argument("--baseline")
    p_validate.add_argument("--previous-revision", required=True)
    p_validate.add_argument("--target-revision", required=True)
    p_validate.add_argument("--output", required=True)

    args = parser.parse_args()
    if args.command == "snapshot":
        result = snapshot_workbook(Path(args.workbook))
        Path(args.output).write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
        return 0

    result = validate(
        Path(args.workbook),
        Path(args.baseline) if args.baseline else None,
        args.previous_revision,
        args.target_revision,
    )
    Path(args.output).write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0 if result["status"] == "VALIDADO" else 1


if __name__ == "__main__":
    raise SystemExit(main())
